from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src=Path('output/report_updated/Email_Spam_Detection_Updated_Report.docx')
d=Document(src)
replacements={
'https://www.kaggle.com/datasets/naserabdullahalam/phishing-email-dataset':'https://www.kaggle.com/datasets/naserabdullahalam/phishing-email-dataset?resource=download',
'https://doi.org/10.24432/C5CC84':'https://archive.ics.uci.edu/dataset/228/sms+spam+collection'
}
found=set()
for para in d.paragraphs:
 for link in para._p.findall(qn('w:hyperlink')):
  rel=d.part.rels[link.get(qn('r:id'))]
  old=rel.target_ref
  if old in replacements:
   url=replacements[old];found.add(old)
   rel._target=url
   texts=list(link.iter(qn('w:t')))
   texts[0].text=url
   for t in texts[1:]:t.text=''
   for r in link.findall(qn('w:r')):
    pr=r.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
    size=pr.find(qn('w:sz'))
    if size is None:size=OxmlElement('w:sz');pr.append(size)
    size.set(qn('w:val'),'19')
assert len(found)==2,found
d.core_properties.subject='Academic submission report with dataset references and measured training results'
d.save('output/submission/Email_Spam_Detection_Submission_Report.docx')
print('Both exact dataset reference links added.')
print('Remaining personal detail placeholders:')
for p in d.paragraphs:
 if '[GUIDE NAME]' in p.text or '[HOD NAME]' in p.text or '[Designation]' in p.text:print(p.text)
