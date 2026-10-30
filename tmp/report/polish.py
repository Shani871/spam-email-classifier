from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Document('output/report/Email_Spam_Detection_Academic_Report.docx')
titles=['Dataset statistics','Email category distribution','Module structure','TF IDF configuration parameters','Risk level mapping','Reported model performance','Reported cross validation results','Evaluation metrics','Dataset sample entries','Scam pattern taxonomy']
nums=[1,2,3,4,5,7,8,6,9,10]
for i,t in enumerate(D.tables):
 n=len(t.columns);ratios=([.23,.13,.64] if i==8 else [.27,.29,.44] if i==9 else [.32,.68] if n==2 else [1/n]*n)
 t.autofit=False
 for c,r in zip(t.columns,ratios):c.width=Inches(6.4*r)
 for row in t.rows:
  for c,r in zip(row.cells,ratios):c.width=Inches(6.4*r)
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 # add captions immediately before table
 el=OxmlElement('w:p');t._tbl.addprevious(el)
 from docx.text.paragraph import Paragraph
 p=Paragraph(el,D._body);p.add_run(f'Table {nums[i]}  {titles[i]}').bold=True;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(6)
 # accuracy formula correction
 for row in t.rows:
  for c in row.cells:
   if c.text=='TP + TN / (TP + TN + FP + FN)':c.text='(TP + TN) / (TP + TN + FP + FN)'
D.save('output/report/Email_Spam_Detection_Academic_Report.docx')
