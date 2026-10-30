import pdfplumber,re
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
path='output/report/Email_Spam_Detection_Academic_Report.docx';d=Document(path)
with pdfplumber.open('tmp/report/finalrender/Email_Spam_Detection_Academic_Report.pdf') as pdf:
 pages=[(p.extract_text() or '').splitlines() for p in pdf.pages]
 def locate(text):
  text=re.sub(r'\s+',' ',text).strip()
  for i,lines in enumerate(pages):
   if i<8:continue
   if any(re.sub(r'\s+',' ',line).strip()==text for line in lines):return i+1
  return None
 active=False
 for p in d.paragraphs:
  if p.text=='TABLE OF CONTENTS':active=True;continue
  if p.text=='LIST OF TABLES':active=False
  if active and p.text.strip():
   num=locate(p.text)
   if num:
    p.text=p.text+'\t'+str(num);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 # list tables
 active=False
 for p in d.paragraphs:
  if p.text=='LIST OF TABLES':active=True;continue
  if p.text=='LIST OF ABBREVIATIONS':active=False
  if active and p.text.strip():
   key=re.sub(r'^\d+\s+','',p.text).lower();found=None
   for i,lines in enumerate(pages):
    if i<8:continue
    if any(key in line.lower() and line.startswith('Table ') for line in lines):found=i+1;break
   if found:p.text+='\t'+str(found);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
d.save(path)
print('pages',len(pages))
